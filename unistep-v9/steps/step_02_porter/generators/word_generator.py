"""
═══════════════════════════════════════════════════════════════════
Porter's 5 Forces Word Report Generator
═══════════════════════════════════════════════════════════════════
"""
import logging
from datetime import datetime

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

logger = logging.getLogger(__name__)


# ═══ Colors ═══
INK = RGBColor(0x0F, 0x1B, 0x2D)
GOLD = RGBColor(0xB8, 0x89, 0x3A)
GREEN = RGBColor(0x0B, 0x7C, 0x4D)
RED = RGBColor(0xA8, 0x4A, 0x2D)
SLATE = RGBColor(0x3A, 0x4A, 0x5C)
MUTED = RGBColor(0x6B, 0x7B, 0x8E)
PLUM = RGBColor(0x5C, 0x2C, 0x5F)


# ═══════════════════════════════════════════════════════════════════
# Helpers (نفس اللي في PESTLE generator)
# ═══════════════════════════════════════════════════════════════════
def _set_rtl(doc):
    for section in doc.sections:
        sectPr = section._sectPr
        bidi = OxmlElement("w:bidi")
        sectPr.append(bidi)


def _set_paragraph_rtl(p):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    pPr.append(bidi)


def _set_arabic_font(run, font_name="Cairo"):
    run.font.name = font_name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:cs"), font_name)


def _set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def _add_heading(doc, text, level=1, color=None, size=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_paragraph_rtl(p)
    run = p.add_run(text)
    _set_arabic_font(run)
    run.bold = True
    sizes = {1: 24, 2: 18, 3: 15}
    run.font.size = Pt(size or sizes.get(level, 14))
    if color is None:
        color = INK if level == 1 else GOLD if level == 2 else SLATE
    run.font.color.rgb = color
    return p


def _add_para(doc, text, size=11, color=None, bold=False, italic=False, align="right"):
    p = doc.add_paragraph()
    if align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_rtl(p)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    _set_arabic_font(run)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    return p


def _add_bullet(doc, text, size=11):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_paragraph_rtl(p)
    p.paragraph_format.right_indent = Cm(0.5)
    run = p.add_run(f"◆  {text}")
    _set_arabic_font(run)
    run.font.size = Pt(size)
    return p


def _threat_emoji(level: str) -> str:
    return {"high": "🔴", "medium": "🟡", "low": "🟢"}.get(level, "⚪")


def _threat_color(level: str) -> RGBColor:
    return {"high": RED, "medium": GOLD, "low": GREEN}.get(level, MUTED)


def _force_label_ar(force: str) -> str:
    return {
        "direct_competitors": "حدّة المنافسة الحالية",
        "substitutes": "تهديد البدائل",
        "new_entrants": "تهديد الوافدين الجدد",
        "donor_power": "القوة التفاوضية للممولين",
        "supplier_power": "القوة التفاوضية للكوادر",
    }.get(force, force)


# ═══════════════════════════════════════════════════════════════════
# Section builders
# ═══════════════════════════════════════════════════════════════════
def _add_cover(doc, client_name):
    for _ in range(8):
        doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_rtl(p)
    run = p.add_run("Porter's 5 Forces")
    _set_arabic_font(run)
    run.bold = True
    run.font.size = Pt(32)
    run.font.color.rgb = INK
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_rtl(p)
    run = p.add_run("تحليل القوى التنافسية الخمس")
    _set_arabic_font(run)
    run.font.size = Pt(20)
    run.font.color.rgb = GOLD
    
    for _ in range(3):
        doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_rtl(p)
    run = p.add_run(client_name)
    _set_arabic_font(run)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = SLATE
    
    for _ in range(8):
        doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_rtl(p)
    run = p.add_run(f"UNISTEP — Step 2/12  •  {datetime.now().strftime('%B %Y')}")
    _set_arabic_font(run)
    run.font.size = Pt(11)
    run.font.color.rgb = MUTED
    run.italic = True
    
    doc.add_page_break()


def _add_competitive_position_section(doc, position_data):
    """قسم التقييم التنافسي العام"""
    _add_heading(doc, "التموضع التنافسي العام", level=1)
    
    pos = position_data.get("competitive_position", "average")
    pos_ar = {"strong": "🟢 قوي", "average": "🟡 متوسط", "weak": "🔴 ضعيف"}.get(pos, pos)
    pos_color = {"strong": GREEN, "average": GOLD, "weak": RED}.get(pos, MUTED)
    
    attr = position_data.get("overall_attractiveness", "medium")
    attr_ar = {"high": "🟢 عالية", "medium": "🟡 متوسطة", "low": "🔴 منخفضة"}.get(attr, attr)
    
    # Two columns table
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    headers = ["التموضع التنافسي", "جاذبية القطاع"]
    values = [pos_ar, attr_ar]
    
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        _set_cell_shading(cell, "0F1B2D")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_paragraph_rtl(p)
        run = p.add_run(h)
        _set_arabic_font(run)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    for i, v in enumerate(values):
        cell = table.rows[1].cells[i]
        _set_cell_shading(cell, "F2EEE5")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_paragraph_rtl(p)
        run = p.add_run(v)
        _set_arabic_font(run)
        run.bold = True
        run.font.size = Pt(14)
    
    doc.add_paragraph()
    
    # Executive summary
    _add_heading(doc, "الملخص التنفيذي", level=2)
    _add_para(doc, position_data.get("executive_summary", "—"), size=12)
    doc.add_paragraph()
    
    # Strongest & weakest forces
    sf = position_data.get("strongest_force", {})
    wf = position_data.get("weakest_force", {})
    
    _add_heading(doc, "🔴 أقوى تهديد ضد الجمعية", level=2, color=RED)
    _add_para(doc, sf.get("force_label_ar", "—"), size=14, bold=True, color=RED)
    _add_para(doc, sf.get("rationale", ""))
    doc.add_paragraph()
    
    _add_heading(doc, "🟢 أضعف تهديد (فرصة)", level=2, color=GREEN)
    _add_para(doc, wf.get("force_label_ar", "—"), size=14, bold=True, color=GREEN)
    _add_para(doc, wf.get("rationale", ""))
    doc.add_paragraph()
    
    # Strategic recommendations
    recs = position_data.get("strategic_recommendations", [])
    if recs:
        _add_heading(doc, "💡 التوصيات الاستراتيجية", level=2)
        for r in recs:
            _add_para(doc, f"{r.get('priority', '—')}. {r.get('title', '')}",
                      size=13, bold=True, color=GOLD)
            _add_para(doc, r.get("description", ""))
            addresses = r.get("addresses_force", "")
            if addresses:
                _add_para(doc, f"تعالج قوة: {_force_label_ar(addresses)}",
                          size=10, italic=True, color=MUTED)
            doc.add_paragraph()


def _add_threat_summary_table(doc, threat_summary):
    """جدول ملخص الـ threat levels للقوى الخمس"""
    _add_heading(doc, "ملخص مستويات التهديد للقوى الخمس", level=2)
    
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    # Header
    for i, h in enumerate(["القوة", "مستوى التهديد", "الرمز"]):
        cell = table.rows[0].cells[i]
        _set_cell_shading(cell, "0F1B2D")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _set_paragraph_rtl(p)
        run = p.add_run(h)
        _set_arabic_font(run)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    # Sort forces by threat level (high first)
    threat_order = {"high": 0, "medium": 1, "low": 2}
    sorted_forces = sorted(
        threat_summary.items(),
        key=lambda x: threat_order.get(x[1].get("threat_level", "medium"), 99)
    )
    
    for force, info in sorted_forces:
        row = table.add_row()
        for cell in row.cells:
            _set_cell_shading(cell, "FAF8F4")
        
        cells = row.cells
        for i, val in enumerate([
            info.get("label", "—"),
            {"high": "عالي", "medium": "متوسط", "low": "منخفض"}.get(info.get("threat_level"), "—"),
            _threat_emoji(info.get("threat_level", "")),
        ]):
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _set_paragraph_rtl(p)
            run = p.add_run(val)
            _set_arabic_font(run)
            run.font.size = Pt(11)
            if i == 1:
                run.bold = True
                run.font.color.rgb = _threat_color(info.get("threat_level", ""))
            elif i == 2:
                run.font.size = Pt(16)
    
    doc.add_paragraph()


def _add_force_section(doc, force_key, force_data):
    """قسم لقوة واحدة من الـ 5"""
    label = force_data.get("force_label_ar", force_key)
    threat = force_data.get("threat_level", "medium")
    emoji = _threat_emoji(threat)
    
    _add_heading(doc, f"{emoji} {label}", level=2, color=_threat_color(threat))
    
    # Threat level + summary
    _add_para(doc, f"مستوى التهديد: {threat}", bold=True, color=_threat_color(threat))
    _add_para(doc, force_data.get("summary", ""), size=11)
    doc.add_paragraph()
    
    # Factors table
    factors = force_data.get("factors", [])
    if factors:
        _add_heading(doc, "العوامل المؤثرة", level=3)
        
        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
        for i, h in enumerate(["العامل", "الوصف", "التأثير"]):
            cell = table.rows[0].cells[i]
            _set_cell_shading(cell, "0F1B2D")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _set_paragraph_rtl(p)
            run = p.add_run(h)
            _set_arabic_font(run)
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
        for f in factors:
            row = table.add_row().cells
            vals = [f.get("factor", ""), f.get("description", ""),
                    {"high": "عالي", "medium": "متوسط", "low": "منخفض"}.get(f.get("impact"), "—")]
            for i, val in enumerate(vals):
                p = row[i].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                _set_paragraph_rtl(p)
                run = p.add_run(val)
                _set_arabic_font(run)
                run.font.size = Pt(10)
        
        doc.add_paragraph()
    
    # Key insights
    insights = force_data.get("key_insights", [])
    if insights:
        _add_heading(doc, "رؤى رئيسية", level=3, color=GOLD)
        for ins in insights:
            _add_bullet(doc, ins)
        doc.add_paragraph()
    
    # Strategic implications
    impl = force_data.get("strategic_implications", "")
    if impl:
        _add_heading(doc, "التضمينات الاستراتيجية", level=3, color=PLUM)
        _add_para(doc, impl, size=11)
        doc.add_paragraph()
    
    doc.add_page_break()


# ═══════════════════════════════════════════════════════════════════
# Main entry
# ═══════════════════════════════════════════════════════════════════
def generate_porter_report(workflow_result: dict, output_path: str) -> str:
    """توليد تقرير Word لـ Porter's 5 Forces"""
    doc = Document()
    _set_rtl(doc)
    
    style = doc.styles["Normal"]
    style.font.name = "Cairo"
    style.font.size = Pt(11)
    
    client_info = workflow_result.get("client", {})
    client_name = client_info.get("name", "غير محدد")
    forces = workflow_result.get("forces", {})
    threat_summary = workflow_result.get("threat_summary", {})
    position = workflow_result.get("competitive_position", {})
    total_cost = workflow_result.get("total_cost_usd", 0)
    
    # Cover
    _add_cover(doc, client_name)
    
    # Competitive position (synthesis)
    _add_competitive_position_section(doc, position)
    
    doc.add_page_break()
    
    # Threat summary table
    _add_threat_summary_table(doc, threat_summary)
    doc.add_page_break()
    
    # Detailed forces
    _add_heading(doc, "التحليل التفصيلي للقوى الخمس", level=1)
    doc.add_paragraph()
    
    force_order = ["direct_competitors", "substitutes", "new_entrants",
                   "donor_power", "supplier_power"]
    for force_key in force_order:
        if force_key in forces:
            _add_force_section(doc, force_key, forces[force_key])
    
    # Footer
    _add_para(doc, "—" * 40, align="center", color=MUTED)
    _add_para(
        doc,
        f"UNISTEP Platform • Step 2/12 (Porter) • تكلفة التحليل: ${total_cost:.4f} • {datetime.now().strftime('%Y-%m-%d')}",
        align="center", color=MUTED, size=9, italic=True,
    )
    
    doc.save(output_path)
    logger.info(f"✅ Porter Word report saved: {output_path}")
    return output_path
