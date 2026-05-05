"""
═══════════════════════════════════════════════════════════════════
PESTLE Word Report Generator
═══════════════════════════════════════════════════════════════════

ينتج تقرير Word احترافي:
- صفحة الغلاف
- ملخص تنفيذي
- 6 أقسام (واحد لكل عامل)
- ملخص الفرص والتهديدات الرئيسية
- توصيات استراتيجية
"""
import logging
from datetime import datetime

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

logger = logging.getLogger(__name__)


# ═══ Colors palette ═══
INK = RGBColor(0x0F, 0x1B, 0x2D)
GOLD = RGBColor(0xB8, 0x89, 0x3A)
GREEN = RGBColor(0x0B, 0x7C, 0x4D)
RED = RGBColor(0xA8, 0x4A, 0x2D)
SLATE = RGBColor(0x3A, 0x4A, 0x5C)
MUTED = RGBColor(0x6B, 0x7B, 0x8E)


# ═══════════════════════════════════════════════════════════════════
# Helpers
# ═══════════════════════════════════════════════════════════════════
def _set_rtl(doc):
    """RTL للوثيقة كاملة"""
    for section in doc.sections:
        sectPr = section._sectPr
        bidi = OxmlElement("w:bidi")
        sectPr.append(bidi)


def _set_paragraph_rtl(p):
    """RTL للفقرة"""
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    pPr.append(bidi)


def _set_arabic_font(run, font_name="Cairo"):
    run.font.name = font_name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:cs"), font_name)


def _set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def _add_heading(doc, text, level=1, color=None, size=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_paragraph_rtl(p)
    run = p.add_run(text)
    _set_arabic_font(run, "Cairo")
    run.bold = True
    
    sizes = {1: 24, 2: 18, 3: 15}
    run.font.size = Pt(size or sizes.get(level, 14))
    
    if color is None:
        color = INK if level == 1 else GOLD if level == 2 else SLATE
    run.font.color.rgb = color
    
    return p


def _add_para(doc, text, size=11, color=None, bold=False, italic=False, align="right"):
    p = doc.add_paragraph()
    if align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_rtl(p)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    _set_arabic_font(run)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    return p


def _add_bullet(doc, text, bullet_char="◆", size=11):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_paragraph_rtl(p)
    p.paragraph_format.right_indent = Cm(0.5)
    run = p.add_run(f"{bullet_char}  {text}")
    _set_arabic_font(run)
    run.font.size = Pt(size)
    return p


# ═══════════════════════════════════════════════════════════════════
# Section builders
# ═══════════════════════════════════════════════════════════════════
def _add_cover_page(doc, client_name, total_cost):
    """صفحة الغلاف"""
    # عنوان كبير في المنتصف
    for _ in range(8):
        doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_rtl(p)
    run = p.add_run("تحليل PESTLE")
    _set_arabic_font(run)
    run.bold = True
    run.font.size = Pt(36)
    run.font.color.rgb = INK
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_rtl(p)
    run = p.add_run("مسح البيئة الكلية")
    _set_arabic_font(run)
    run.font.size = Pt(20)
    run.font.color.rgb = GOLD
    
    for _ in range(3):
        doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_rtl(p)
    run = p.add_run(client_name)
    _set_arabic_font(run)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = SLATE
    
    for _ in range(8):
        doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_rtl(p)
    run = p.add_run(f"UNISTEP — Step 1/12  •  {datetime.now().strftime('%B %Y')}")
    _set_arabic_font(run)
    run.font.size = Pt(11)
    run.font.color.rgb = MUTED
    run.italic = True
    
    doc.add_page_break()


def _add_executive_summary(doc, synthesis):
    """قسم الملخص التنفيذي"""
    _add_heading(doc, "الملخص التنفيذي", level=1)
    _add_para(doc, synthesis.get("executive_summary", "—"), size=12)
    doc.add_paragraph()
    
    # Strategic posture
    posture = synthesis.get("strategic_posture", "balanced")
    posture_ar = {
        "defensive": "🛡️ دفاعي",
        "offensive": "⚔️ هجومي",
        "balanced": "⚖️ متوازن",
    }.get(posture, posture)
    
    _add_heading(doc, "التموضع الاستراتيجي المُوصى به", level=2)
    _add_para(doc, posture_ar, size=14, bold=True, color=GOLD)
    _add_para(doc, synthesis.get("posture_rationale", ""), size=11, italic=True)
    doc.add_paragraph()


def _add_top_items(doc, items, title, color, severity_field="priority"):
    """قسم لأهم الفرص أو التهديدات"""
    _add_heading(doc, title, level=2)
    
    if not items:
        _add_para(doc, "لا توجد عناصر مُحددة.", italic=True, color=MUTED)
        return
    
    for i, item in enumerate(items, 1):
        title_text = item.get("title", "—")
        desc = item.get("description", "")
        factors = item.get("factors", [])
        
        _add_para(doc, f"{i}. {title_text}", size=13, bold=True, color=color)
        _add_para(doc, desc, size=11)
        
        if factors:
            factors_ar = [_factor_ar(f) for f in factors]
            _add_para(doc, f"العوامل المرتبطة: {' • '.join(factors_ar)}",
                      size=10, italic=True, color=MUTED)
        doc.add_paragraph()


def _factor_ar(factor: str) -> str:
    return {
        "political": "السياسية",
        "economic": "الاقتصادية",
        "social": "الاجتماعية",
        "technology": "التقنية",
        "legal": "القانونية",
        "environmental": "البيئية",
    }.get(factor, factor)


def _add_factor_section(doc, factor_key, factor_data):
    """قسم لعامل واحد من الـ 6"""
    factor_label = factor_data.get("factor_label_ar", factor_key)
    impact = factor_data.get("impact_level", "medium")
    impact_emoji = {"high": "🔴", "medium": "🟡", "low": "🟢"}.get(impact, "⚪")
    
    _add_heading(doc, f"العوامل {factor_label}", level=2)
    
    # Impact + summary
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_paragraph_rtl(p)
    run = p.add_run(f"{impact_emoji} مستوى التأثير: ")
    _set_arabic_font(run)
    run.bold = True
    run.font.size = Pt(11)
    run = p.add_run({"high": "عالي", "medium": "متوسط", "low": "منخفض"}.get(impact, impact))
    _set_arabic_font(run)
    run.font.size = Pt(11)
    run.font.color.rgb = {"high": RED, "medium": GOLD, "low": GREEN}.get(impact, MUTED)
    
    _add_para(doc, factor_data.get("summary", ""), size=11)
    doc.add_paragraph()
    
    # Drivers table
    drivers = factor_data.get("key_drivers", [])
    if drivers:
        _add_heading(doc, "العوامل الرئيسية", level=3)
        
        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
        # Header
        hdr = table.rows[0].cells
        for i, h in enumerate(["العامل", "الوصف", "التأثير"]):
            cell = hdr[i]
            _set_cell_shading(cell, "0F1B2D")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _set_paragraph_rtl(p)
            run = p.add_run(h)
            _set_arabic_font(run)
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
        # Rows
        for d in drivers:
            row = table.add_row().cells
            row[0].text = ""
            row[1].text = ""
            row[2].text = ""
            for i, val in enumerate([d.get("name", ""), d.get("description", ""),
                                     {"high": "عالي", "medium": "متوسط", "low": "منخفض"}.get(d.get("impact"), "—")]):
                cell = row[i]
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                _set_paragraph_rtl(p)
                run = p.add_run(val)
                _set_arabic_font(run)
                run.font.size = Pt(10)
        
        doc.add_paragraph()
    
    # Opportunities
    opps = factor_data.get("opportunities", [])
    if opps:
        _add_heading(doc, "الفرص", level=3, color=GREEN)
        for o in opps:
            _add_bullet(doc, f"{o.get('title', '')}: {o.get('description', '')}")
        doc.add_paragraph()
    
    # Threats
    threats = factor_data.get("threats", [])
    if threats:
        _add_heading(doc, "التهديدات", level=3, color=RED)
        for t in threats:
            _add_bullet(doc, f"{t.get('title', '')}: {t.get('description', '')}")
        doc.add_paragraph()
    
    # Extra (مثل barriers_to_entry للـ legal)
    extra = factor_data.get("extra", {})
    if "barriers_to_entry" in extra:
        bte = extra["barriers_to_entry"]
        _add_heading(doc, "حواجز الدخول للقطاع (لتحليل Porter لاحقاً)", level=3, color=GOLD)
        level = bte.get("level", "medium")
        _add_para(doc, f"مستوى الحواجز: {level}", bold=True)
        _add_para(doc, bte.get("description", ""))
        for req in bte.get("key_requirements", []):
            _add_bullet(doc, req)
        doc.add_paragraph()
    
    doc.add_page_break()


# ═══════════════════════════════════════════════════════════════════
# Main entry
# ═══════════════════════════════════════════════════════════════════
def generate_pestle_report(workflow_result: dict, output_path: str) -> str:
    """
    توليد تقرير Word لـ PESTLE.
    
    Args:
        workflow_result: المخرج من PestleStep
        output_path: مسار حفظ الملف
    
    Returns:
        output_path
    """
    doc = Document()
    _set_rtl(doc)
    
    # Set default font
    style = doc.styles["Normal"]
    style.font.name = "Cairo"
    style.font.size = Pt(11)
    
    client_info = workflow_result.get("client", {})
    client_name = client_info.get("name", "غير محدد")
    synthesis = workflow_result.get("synthesis", {})
    factors = workflow_result.get("factors", {})
    total_cost = workflow_result.get("total_cost_usd", 0)
    
    # Cover
    _add_cover_page(doc, client_name, total_cost)
    
    # Executive summary
    _add_executive_summary(doc, synthesis)
    
    # Top opportunities & threats
    _add_top_items(
        doc,
        synthesis.get("top_opportunities", []),
        "🟢 أهم 3 فرص استراتيجية",
        GREEN,
    )
    
    _add_top_items(
        doc,
        synthesis.get("top_threats", []),
        "🔴 أهم 3 تهديدات",
        RED,
    )
    
    # Cross-factor insights
    insights = synthesis.get("cross_factor_insights", [])
    if insights:
        _add_heading(doc, "🔗 رؤى مشتركة بين العوامل", level=2)
        for i in insights:
            connected = " ، ".join([_factor_ar(f) for f in i.get("connected_factors", [])])
            _add_para(doc, f"العوامل المرتبطة: {connected}", size=10, italic=True, color=MUTED)
            _add_para(doc, i.get("insight", ""), size=11)
            doc.add_paragraph()
    
    doc.add_page_break()
    
    # Detailed factor analysis
    _add_heading(doc, "التحليل التفصيلي للعوامل الستة", level=1)
    doc.add_paragraph()
    
    # Order: P-E-S-T-L-E
    factor_order = ["political", "economic", "social", "technology", "legal", "environmental"]
    for factor_key in factor_order:
        if factor_key in factors:
            _add_factor_section(doc, factor_key, factors[factor_key])
    
    # Footer info
    _add_para(doc, "—" * 40, align="center", color=MUTED)
    _add_para(
        doc,
        f"UNISTEP Platform • Step 1/12 (PESTLE) • تكلفة التحليل: ${total_cost:.4f} • {datetime.now().strftime('%Y-%m-%d')}",
        align="center", color=MUTED, size=9, italic=True,
    )
    
    # Save
    doc.save(output_path)
    logger.info(f"✅ PESTLE Word report saved: {output_path}")
    return output_path
