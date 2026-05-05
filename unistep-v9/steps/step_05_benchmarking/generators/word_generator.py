"""
═══════════════════════════════════════════════════════════
Word Generator v7 — التقرير الكامل
═══════════════════════════════════════════════════════════

يولد ملف Word احترافي بالـ 3 مخرجات:
1. مخرج 1: معايير الاختيار + تقييم الجهات
2. مخرج 2: بيانات المنافسين (10 صفوف لكل جهة)
3. مخرج 3: المصفوفة الشاملة (7 أعمدة)

+ صفحة غلاف، بيانات العميل، تحليل، توصيات، Audit log
"""
from datetime import datetime
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


COLOR_DARK = RGBColor(0x1F, 0x29, 0x37)
COLOR_AMBER = RGBColor(0xD9, 0x77, 0x06)
COLOR_GRAY = RGBColor(0x6B, 0x72, 0x80)
COLOR_GREEN = RGBColor(0x05, 0x96, 0x69)
COLOR_RED = RGBColor(0xDC, 0x26, 0x26)
COLOR_BLUE = RGBColor(0x25, 0x63, 0xEB)


SELECTION_CRITERIA_7 = [
    ("تطابق الفئة المستهدفة", "نفس الفئة الدقيقة (مش القطاع العام)", "إلزامي", "target_match"),
    ("نفس التخصص", "تطابق التخصص الفرعي", "إلزامي", "specialization_match"),
    ("تنوع البرامج", "3 مبادرات على الأقل", "إلزامي", "has_diverse_programs"),
    ("النضج المؤسسي", "هيكل + تقارير سنوية", "إلزامي", "institutional_maturity"),
    ("التوافق مع رؤية 2030", "كلمات: تمكين، استدامة، حوكمة", "إلزامي", "vision_2030_aligned"),
    ("التغطية الجغرافية", "تنوع جغرافي", "ثانوي", "geographic_diversity"),
    ("توفر البيانات", "مؤشرات + تمويل منشورة", "ثانوي", "data_available"),
]


FIELD_LABELS = {
    "name": "اسم الجهة",
    "region": "المنطقة الجغرافية",
    "field_of_work": "مجال العمل",
    "target_audience": "الفئة المستهدفة",
    "vision": "الرؤية",
    "mission": "الرسالة",
    "strategic_goals": "الأهداف الاستراتيجية",
    "programs": "البرامج والمبادرات",
    "metrics_results": "المؤشرات والنتائج",
    "funding_sources": "مصادر التمويل",
}


ROW_ORDER = [
    "name", "region", "field_of_work", "target_audience",
    "vision", "mission", "strategic_goals", "programs",
    "metrics_results", "funding_sources",
]


def _set_rtl(doc):
    sectPr = doc.sections[0]._sectPr
    bidi = OxmlElement('w:bidi')
    sectPr.append(bidi)


def _set_paragraph_rtl(p):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)


def _set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def _set_arabic_font(run):
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:cs'), 'Arial')
    rFonts.set(qn('w:hAnsi'), 'Arial')


def _add_heading(doc, text, level=1, color=None):
    sizes = {1: 22, 2: 16, 3: 13}
    p = doc.add_paragraph()
    _set_paragraph_rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(8)
    
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(sizes.get(level, 11))
    run.bold = True
    run.font.color.rgb = color or (COLOR_DARK if level == 1 else COLOR_AMBER)
    _set_arabic_font(run)
    return p


def _add_para(doc, text, font_size=11, bold=False, color=None, alignment=None):
    p = doc.add_paragraph()
    _set_paragraph_rtl(p)
    p.alignment = alignment or WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(6)
    
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(font_size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    _set_arabic_font(run)
    return p


def _add_hr(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:color'), 'FBBF24')
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    bidi = OxmlElement('w:bidiVisual')
    tblPr.append(bidi)
    
    # Header
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        _set_paragraph_rtl(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.name = 'Arial'
        run.font.size = Pt(10)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _set_arabic_font(run)
        _set_cell_shading(cell, '1F2937')
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    
    # Data
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        for c_idx, val in enumerate(row_data):
            if c_idx >= len(headers):
                break
            cell = row.cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            _set_paragraph_rtl(p)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(str(val) if val is not None else "—")
            run.font.name = 'Arial'
            run.font.size = Pt(9)
            _set_arabic_font(run)
            
            if r_idx % 2 == 0:
                _set_cell_shading(cell, 'F5F3F0')
    
    if col_widths:
        for col, w in zip(table.columns, col_widths):
            for cell in col.cells:
                cell.width = Inches(w)
    
    return table


def generate_word_report(workflow_result: dict, output_path: str) -> str:
    """يولّد التقرير الكامل"""
    doc = Document()
    _set_rtl(doc)
    
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
    
    _build_cover(doc, workflow_result)
    doc.add_page_break()
    
    _build_client_info(doc, workflow_result)
    doc.add_page_break()
    
    _build_output_1(doc, workflow_result)
    doc.add_page_break()
    
    _build_output_2(doc, workflow_result)
    doc.add_page_break()
    
    _build_output_3(doc, workflow_result)
    doc.add_page_break()
    
    _build_analysis(doc, workflow_result)
    doc.add_page_break()
    
    _build_audit(doc, workflow_result)
    
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    return str(output)


def _build_cover(doc, result):
    client = result.get("client", {})
    
    for _ in range(6):
        doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("UNISTEP")
    run.font.name = 'Arial'
    run.font.size = Pt(40)
    run.bold = True
    run.font.color.rgb = COLOR_AMBER
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("•  المقارنة المعيارية التنافسية  •")
    run.font.name = 'Arial'
    run.font.size = Pt(11)
    run.font.color.rgb = COLOR_GRAY
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    _set_paragraph_rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(client.get("name", ""))
    run.font.name = 'Arial'
    run.font.size = Pt(24)
    run.bold = True
    run.font.color.rgb = COLOR_DARK
    _set_arabic_font(run)
    
    p = doc.add_paragraph()
    _set_paragraph_rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(client.get("specialization") or client.get("sector", ""))
    run.font.name = 'Arial'
    run.font.size = Pt(13)
    run.font.color.rgb = COLOR_GRAY
    _set_arabic_font(run)
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(datetime.now().strftime("%Y / %m / %d"))
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.font.color.rgb = COLOR_AMBER


def _build_client_info(doc, result):
    client = result.get("client", {})
    
    _add_heading(doc, "📋 بيانات الجمعية العميل", level=1)
    _add_hr(doc)
    
    fields = [
        ("اسم الجمعية", client.get("name", "")),
        ("القطاع العام", client.get("sector", "")),
        ("التخصص الدقيق", client.get("specialization", "")),
        ("الفئة المستهدفة", client.get("target_audience", "")),
        ("حجم الجمعية", client.get("organization_size", "")),
        ("المنطقة", client.get("region", "—")),
    ]
    
    if client.get("strategic_goals"):
        goals = client["strategic_goals"]
        text = "\n".join(f"• {g}" for g in goals) if isinstance(goals, list) else str(goals)
        fields.append(("الأهداف الاستراتيجية", text))
    
    if client.get("vision"):
        fields.append(("الرؤية", client["vision"]))
    if client.get("mission"):
        fields.append(("الرسالة", client["mission"]))
    
    if client.get("funding_sources"):
        fund = client["funding_sources"]
        text = "، ".join(fund) if isinstance(fund, list) else str(fund)
        fields.append(("مصادر التمويل", text))
    
    if client.get("current_kpis"):
        fields.append(("مؤشرات الأداء", client["current_kpis"]))
    
    rows = [[label, value] for label, value in fields]
    _add_table(doc, ["البيان", "التفاصيل"], rows, col_widths=[2, 5])


def _build_output_1(doc, result):
    _add_heading(doc, "📊 المخرج 1: معايير اختيار المقارنات المعيارية", level=1)
    _add_hr(doc)
    
    _add_para(doc,
        "تم اختيار الجهات بناءً على 7 معايير: 5 إلزامية + 2 ثانوية.",
        font_size=11)
    doc.add_paragraph()
    
    # 1.1
    _add_heading(doc, "1.1 — وصف المعايير", level=2)
    rows = []
    for criterion, desc, level, _ in SELECTION_CRITERIA_7:
        icon = "🔴" if level == "إلزامي" else "🟡"
        rows.append([f"{icon} {level}", criterion, desc])
    _add_table(doc, ["المستوى", "المعيار", "الوصف"], rows, col_widths=[1, 2, 4])
    doc.add_paragraph()
    
    # 1.2
    accepted = [
        o for o in result.get("local_organizations", []) + result.get("international_organizations", [])
        if o.get("is_accepted")
    ]
    if accepted:
        _add_heading(doc, "1.2 — تقييم الجهات على المعايير", level=2)
        
        headers = ["الجهة"] + [c[0][:14] for c in SELECTION_CRITERIA_7]
        eval_rows = []
        for org in accepted:
            row = [org.get("name", "")]
            for _, _, _, criterion_id in SELECTION_CRITERIA_7:
                val = org.get(criterion_id, False)
                row.append("✅" if val else "—")
            eval_rows.append(row)
        _add_table(doc, headers, eval_rows)


def _build_output_2(doc, result):
    accepted = [
        o for o in result.get("local_organizations", []) + result.get("international_organizations", [])
        if o.get("is_accepted")
    ]
    
    _add_heading(doc, "📊 المخرج 2: بيانات المنافسين", level=1)
    _add_hr(doc)
    
    _add_para(doc, f"عدد الجهات المنافسة: {len(accepted)}", font_size=11)
    doc.add_paragraph()
    
    local_orgs = [o for o in accepted if o.get("type") == "local"]
    intl_orgs = [o for o in accepted if o.get("type") == "international"]
    
    if local_orgs:
        _add_heading(doc, "أولاً — الجهات المحلية", level=2)
        for i, org in enumerate(local_orgs, 1):
            _build_org_table(doc, org, i)
            doc.add_paragraph()
    
    if intl_orgs:
        _add_heading(doc, "ثانياً — الجهات الدولية", level=2)
        for i, org in enumerate(intl_orgs, 1):
            _build_org_table(doc, org, i)
            doc.add_paragraph()


def _build_org_table(doc, org, index):
    name = org.get("name", "")
    
    p = doc.add_paragraph()
    _set_paragraph_rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(f"{index}. {name}")
    run.font.name = 'Arial'
    run.font.size = Pt(14)
    run.bold = True
    run.font.color.rgb = COLOR_AMBER
    _set_arabic_font(run)
    
    _add_para(doc,
        f"🔗 {org.get('url', '')}   |   📍 {org.get('country', '')}",
        font_size=9, color=COLOR_GRAY)
    
    fields_dict = {f.get("name"): f for f in org.get("fields", [])}
    rows = []
    for fname in ROW_ORDER:
        label = FIELD_LABELS.get(fname, fname)
        f = fields_dict.get(fname)
        value = f.get("value") if f and f.get("value") else "غير متاح"
        rows.append([label, str(value)])
    _add_table(doc, ["البيان", "التفاصيل"], rows, col_widths=[1.7, 5.0])


def _build_output_3(doc, result):
    matrix = result.get("initiative_matrix", [])
    
    _add_heading(doc, "📊 المخرج 3: المصفوفة الشاملة", level=1)
    _add_hr(doc)
    
    _add_para(doc,
        f"تم فلترة المبادرات على تخصصك الدقيق. إجمالي المبادرات المتطابقة: {len(matrix)}.",
        font_size=11)
    doc.add_paragraph()
    
    local_matrix = [m for m in matrix if m.get("source_type") == "local"]
    intl_matrix = [m for m in matrix if m.get("source_type") == "international"]
    
    headers = [
        "اسم الجمعية", "اسم المبادرة", "وصف المبادرة",
        "الهدف الاستراتيجي", "الفئة المستهدفة",
        "نوع التمويل", "مؤشرات الأداء"
    ]
    col_widths = [1.4, 1.5, 1.8, 1.4, 1.2, 1.2, 1.5]
    
    if local_matrix:
        _add_heading(doc, f"3.1 — المصفوفة المحلية ({len(local_matrix)} مبادرة)", level=2)
        rows = []
        for init in local_matrix:
            metric = init.get("metric") or "—"
            if init.get("results"):
                metric += f" | {init.get('results')}"
            rows.append([
                init.get("source_org", ""),
                init.get("name", ""),
                (init.get("description") or "—")[:200],
                (init.get("objective") or "—")[:120],
                (init.get("target_beneficiaries") or "—")[:100],
                (init.get("funding_type") or "—")[:80],
                metric[:120],
            ])
        _add_table(doc, headers, rows, col_widths=col_widths)
        doc.add_paragraph()
    
    if intl_matrix:
        _add_heading(doc, f"3.2 — المصفوفة الدولية ({len(intl_matrix)} مبادرة)", level=2)
        rows = []
        for init in intl_matrix:
            metric = init.get("metric") or "—"
            if init.get("results"):
                metric += f" | {init.get('results')}"
            rows.append([
                init.get("source_org", ""),
                init.get("name", ""),
                (init.get("description") or "—")[:200],
                (init.get("objective") or "—")[:120],
                (init.get("target_beneficiaries") or "—")[:100],
                (init.get("funding_type") or "—")[:80],
                metric[:120],
            ])
        _add_table(doc, headers, rows, col_widths=col_widths)


def _build_analysis(doc, result):
    analysis = result.get("analysis", {})
    
    # Strategic Insights
    insights = analysis.get("strategic_insights", [])
    if insights:
        _add_heading(doc, "💡 الرؤى الاستراتيجية", level=1)
        _add_hr(doc)
        for i, insight in enumerate(insights, 1):
            p = doc.add_paragraph()
            _set_paragraph_rtl(p)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.right_indent = Cm(0.5)
            
            run = p.add_run(f"{i}.  ")
            run.font.size = Pt(13)
            run.bold = True
            run.font.color.rgb = COLOR_AMBER
            
            run = p.add_run(insight)
            run.font.name = 'Arial'
            run.font.size = Pt(11)
            run.font.color.rgb = COLOR_DARK
            _set_arabic_font(run)
        doc.add_paragraph()
    
    # Recommendations
    recs = analysis.get("top_3_recommendations", [])
    if recs:
        _add_heading(doc, "🎯 التوصيات الاستراتيجية", level=1)
        _add_hr(doc)
        rows = []
        for rec in recs:
            rows.append([
                f"#{rec.get('priority', '')}",
                rec.get("action", ""),
                rec.get("expected_impact", ""),
                rec.get("timeframe", ""),
            ])
        _add_table(doc, ["الأولوية", "الإجراء", "الأثر", "الإطار الزمني"], rows,
                   col_widths=[0.7, 3, 1.8, 1.3])
        doc.add_paragraph()
    
    # Selected Initiatives
    selected = analysis.get("selected_initiatives", [])
    if selected:
        _add_heading(doc, "⭐ المبادرات المختارة + الفرص", level=1)
        _add_hr(doc)
        
        for i, init in enumerate(selected, 1):
            p = doc.add_paragraph()
            _set_paragraph_rtl(p)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.space_before = Pt(8)
            
            run = p.add_run(f"{i}. {init.get('name', '')}")
            run.font.size = Pt(13)
            run.bold = True
            run.font.color.rgb = COLOR_DARK
            _set_arabic_font(run)
            
            run = p.add_run(f"   [{init.get('opportunity_type', '')}]   ")
            run.font.size = Pt(10)
            run.font.color.rgb = COLOR_AMBER
            run.bold = True
            
            run = p.add_run(f"{init.get('applicability_score', 0)}/10")
            run.font.size = Pt(10)
            run.bold = True
            run.font.color.rgb = COLOR_GREEN
            
            details = [
                ("المصدر", init.get("source_org", "")),
                ("ماذا تفعل؟", init.get("what_it_does", "")),
                ("ما حققته؟", init.get("achievements", "—")),
                ("التمويل", init.get("funding_source", "—")),
                ("لماذا لك؟", init.get("rationale", "")),
            ]
            
            for label, value in details:
                p = doc.add_paragraph()
                _set_paragraph_rtl(p)
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.right_indent = Cm(0.5)
                p.paragraph_format.space_after = Pt(2)
                
                run = p.add_run(f"• {label}: ")
                run.font.size = Pt(10)
                run.bold = True
                run.font.color.rgb = COLOR_AMBER
                _set_arabic_font(run)
                
                run = p.add_run(str(value) if value else "—")
                run.font.size = Pt(10)
                _set_arabic_font(run)
    
    # Gap Analysis
    gap = analysis.get("gap_analysis", {})
    if gap:
        doc.add_paragraph()
        _add_heading(doc, "🔍 تحليل الفجوات", level=1)
        _add_hr(doc)
        
        sections = [
            ("❌ برامج ناقصة", gap.get("missing_programs", []), COLOR_RED),
            ("⚠️ مجالات تحتاج تطويراً", gap.get("weak_areas", []), COLOR_AMBER),
            ("✅ مزايا تنافسية", gap.get("competitive_advantages", []), COLOR_GREEN),
        ]
        for title, items, color in sections:
            _add_heading(doc, title, level=2)
            if not items:
                _add_para(doc, "(لا توجد بيانات)", font_size=10, color=COLOR_GRAY)
            else:
                for item in items:
                    p = doc.add_paragraph()
                    _set_paragraph_rtl(p)
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    p.paragraph_format.left_indent = Cm(0.5)
                    p.paragraph_format.right_indent = Cm(0.5)
                    p.paragraph_format.space_after = Pt(4)
                    
                    run = p.add_run("●  ")
                    run.font.size = Pt(10)
                    run.font.color.rgb = color
                    run.bold = True
                    
                    run = p.add_run(item)
                    run.font.size = Pt(10)
                    _set_arabic_font(run)
            doc.add_paragraph()


def _build_audit(doc, result):
    accepted = [
        o for o in result.get("local_organizations", []) + result.get("international_organizations", [])
        if o.get("is_accepted")
    ]
    
    _add_heading(doc, "🔒 سجل التحقق", level=1)
    _add_hr(doc)
    
    _add_para(doc,
        "كل بيان في هذا التقرير له اقتباس حرفي تم التحقق من وجوده فعلياً في الموقع المصدر.",
        font_size=10, color=COLOR_GREEN)
    doc.add_paragraph()
    
    rows = []
    for org in accepted:
        for f in org.get("fields", []):
            if not f.get("passed"):
                continue
            label = FIELD_LABELS.get(f.get("name", ""), f.get("name", ""))
            rows.append([
                org.get("name", "")[:25],
                label,
                str(f.get("value", ""))[:80],
                str(f.get("exact_quote", ""))[:150],
            ])
    
    if rows:
        _add_table(doc,
            ["الجمعية", "البيان", "القيمة", "الاقتباس الحرفي"],
            rows, col_widths=[1.5, 1.3, 2.0, 2.5])
