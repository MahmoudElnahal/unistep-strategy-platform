"""
═══════════════════════════════════════════════════════════
Excel Generator v7
═══════════════════════════════════════════════════════════

ينتج 4 ملفات:
1. مخرج_1_معايير_الاختيار.xlsx
2. مخرج_2_بيانات_المنافسين.xlsx
3. مخرج_3_المصفوفة.xlsx
4. التقرير_الموحد.xlsx
"""
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


COLORS = {
    "amber": "FBBF24", "amber_light": "FEF3C7",
    "cream": "F5F3F0", "dark": "1F2937",
    "green_light": "D1FAE5", "red_light": "FEE2E2", "blue_light": "DBEAFE",
}


SELECTION_CRITERIA_7 = [
    ("تطابق الفئة المستهدفة", "نفس الفئة الدقيقة", "إلزامي", "target_match"),
    ("نفس التخصص", "تطابق التخصص الفرعي", "إلزامي", "specialization_match"),
    ("تنوع البرامج", "3 مبادرات على الأقل", "إلزامي", "has_diverse_programs"),
    ("النضج المؤسسي", "هيكل + تقارير سنوية", "إلزامي", "institutional_maturity"),
    ("التوافق مع رؤية 2030", "كلمات: تمكين، استدامة، حوكمة", "إلزامي", "vision_2030_aligned"),
    ("التغطية الجغرافية", "تنوع جغرافي", "ثانوي", "geographic_diversity"),
    ("توفر البيانات", "مؤشرات + تمويل", "ثانوي", "data_available"),
]


FIELD_LABELS = {
    "name": "اسم الجهة",
    "region": "المنطقة الجغرافية",
    "field_of_work": "مجال العمل",
    "target_audience": "الفئة المستهدفة",
    "vision": "الرؤية",
    "mission": "الرسالة",
    "strategic_goals": "الأهداف الاستراتيجية",
    "programs": "البرامج والمبادرات",
    "metrics_results": "المؤشرات والنتائج",
    "funding_sources": "مصادر التمويل",
}


ROW_ORDER = [
    "name", "region", "field_of_work", "target_audience",
    "vision", "mission", "strategic_goals", "programs",
    "metrics_results", "funding_sources",
]


def main_title_style():
    return {
        "font": Font(name="IBM Plex Sans Arabic", size=16, bold=True, color="FFFFFF"),
        "fill": PatternFill("solid", fgColor=COLORS["dark"]),
        "alignment": Alignment(horizontal="center", vertical="center", wrap_text=True),
    }


def header_style():
    return {
        "font": Font(name="IBM Plex Sans Arabic", size=11, bold=True, color="FFFFFF"),
        "fill": PatternFill("solid", fgColor=COLORS["dark"]),
        "alignment": Alignment(horizontal="center", vertical="center", wrap_text=True),
    }


def label_style():
    return {
        "font": Font(name="IBM Plex Sans Arabic", size=11, bold=True, color=COLORS["dark"]),
        "fill": PatternFill("solid", fgColor=COLORS["amber_light"]),
        "alignment": Alignment(horizontal="right", vertical="center", wrap_text=True),
    }


def value_style():
    return {
        "font": Font(name="IBM Plex Sans Arabic", size=10, color=COLORS["dark"]),
        "alignment": Alignment(horizontal="right", vertical="center", wrap_text=True),
    }


def thin_border():
    side = Side(border_style="thin", color="D1D5DB")
    return Border(top=side, bottom=side, left=side, right=side)


def apply_style(cell, style):
    if "font" in style: cell.font = style["font"]
    if "fill" in style: cell.fill = style["fill"]
    if "alignment" in style: cell.alignment = style["alignment"]
    cell.border = thin_border()


def generate_output_1(workflow_result, output_path):
    """مخرج 1: معايير الاختيار"""
    wb = Workbook()
    wb.remove(wb.active)
    
    accepted = [o for o in workflow_result.get("local_organizations", []) +
                workflow_result.get("international_organizations", []) if o.get("is_accepted")]
    
    # Sheet 1: وصف
    ws1 = wb.create_sheet("وصف المعايير")
    ws1.sheet_view.rightToLeft = True
    
    ws1.merge_cells("A1:C1")
    ws1["A1"] = "جدول معايير اختيار المقارنات المعيارية"
    apply_style(ws1["A1"], main_title_style())
    ws1.row_dimensions[1].height = 40
    
    for col, h in enumerate(["المستوى", "المعيار", "الوصف"], 1):
        apply_style(ws1.cell(row=3, column=col, value=h), header_style())
    ws1.row_dimensions[3].height = 30
    
    for i, (crit, desc, lvl, _) in enumerate(SELECTION_CRITERIA_7, 4):
        icon = "🔴" if lvl == "إلزامي" else "🟡"
        ws1.cell(row=i, column=1, value=f"{icon} {lvl}")
        ws1.cell(row=i, column=2, value=crit)
        ws1.cell(row=i, column=3, value=desc)
        for c in range(1, 4):
            apply_style(ws1.cell(row=i, column=c), value_style())
    
    ws1.column_dimensions["A"].width = 14
    ws1.column_dimensions["B"].width = 30
    ws1.column_dimensions["C"].width = 50
    
    # Sheet 2: تقييم
    if accepted:
        ws2 = wb.create_sheet("تقييم الجهات")
        ws2.sheet_view.rightToLeft = True
        
        ws2.merge_cells("A1:H1")
        ws2["A1"] = "تقييم الجهات المختارة"
        apply_style(ws2["A1"], main_title_style())
        ws2.row_dimensions[1].height = 40
        
        headers = ["الجهة"] + [c[0][:20] for c in SELECTION_CRITERIA_7]
        for col, h in enumerate(headers, 1):
            apply_style(ws2.cell(row=3, column=col, value=h), header_style())
        ws2.row_dimensions[3].height = 50
        
        for i, org in enumerate(accepted, 4):
            ws2.cell(row=i, column=1, value=org.get("name", ""))
            for j, (_, _, _, cid) in enumerate(SELECTION_CRITERIA_7, 2):
                val = org.get(cid, False)
                ws2.cell(row=i, column=j, value="✅" if val else "—")
            for c in range(1, len(headers) + 1):
                apply_style(ws2.cell(row=i, column=c), value_style())
        
        ws2.column_dimensions["A"].width = 30
        for col in range(2, 9):
            ws2.column_dimensions[chr(64 + col)].width = 14
    
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output))
    return str(output)


def generate_output_2(workflow_result, output_path):
    """مخرج 2: بيانات المنافسين"""
    wb = Workbook()
    wb.remove(wb.active)
    
    accepted = [o for o in workflow_result.get("local_organizations", []) +
                workflow_result.get("international_organizations", []) if o.get("is_accepted")]
    
    # Index
    ws_index = wb.create_sheet("الفهرس")
    ws_index.sheet_view.rightToLeft = True
    
    ws_index.merge_cells("A1:D1")
    ws_index["A1"] = "بيانات المنافسين — الفهرس"
    apply_style(ws_index["A1"], main_title_style())
    ws_index.row_dimensions[1].height = 40
    
    for col, h in enumerate(["#", "اسم الجهة", "النوع", "الموقع"], 1):
        apply_style(ws_index.cell(row=3, column=col, value=h), header_style())
    ws_index.row_dimensions[3].height = 30
    
    for i, org in enumerate(accepted, 4):
        type_label = "محلية" if org.get("type") == "local" else "دولية"
        ws_index.cell(row=i, column=1, value=i - 3)
        ws_index.cell(row=i, column=2, value=org.get("name", ""))
        ws_index.cell(row=i, column=3, value=type_label)
        ws_index.cell(row=i, column=4, value=org.get("url", ""))
        for c in range(1, 5):
            apply_style(ws_index.cell(row=i, column=c), value_style())
    
    ws_index.column_dimensions["A"].width = 5
    ws_index.column_dimensions["B"].width = 35
    ws_index.column_dimensions["C"].width = 12
    ws_index.column_dimensions["D"].width = 40
    
    # Sheet لكل جهة
    for i, org in enumerate(accepted, 1):
        sheet_name = f"{i}. {org.get('name', '')[:25]}"
        ws = wb.create_sheet(sheet_name[:31])
        ws.sheet_view.rightToLeft = True
        
        ws.merge_cells("A1:B1")
        ws["A1"] = org.get("name", "")
        apply_style(ws["A1"], main_title_style())
        ws.row_dimensions[1].height = 40
        
        ws.merge_cells("A2:B2")
        ws["A2"] = f"{org.get('url', '')}  |  {org.get('country', '')}"
        apply_style(ws["A2"], {
            "font": Font(name="IBM Plex Sans Arabic", size=9, italic=True),
            "fill": PatternFill("solid", fgColor=COLORS["cream"]),
            "alignment": Alignment(horizontal="center", vertical="center"),
        })
        
        apply_style(ws.cell(row=4, column=1, value="اسم البيان"), header_style())
        apply_style(ws.cell(row=4, column=2, value="التفاصيل"), header_style())
        ws.row_dimensions[4].height = 30
        
        fields_dict = {f.get("name"): f for f in org.get("fields", [])}
        row = 5
        for fname in ROW_ORDER:
            label = FIELD_LABELS.get(fname, fname)
            f = fields_dict.get(fname)
            value = f.get("value") if f and f.get("value") else "غير متاح"
            
            apply_style(ws.cell(row=row, column=1, value=label), label_style())
            apply_style(ws.cell(row=row, column=2, value=str(value)), value_style())
            ws.row_dimensions[row].height = min(180, max(35, len(str(value)) // 5))
            row += 1
        
        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 75
    
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output))
    return str(output)


def _format_funding_summary(init: dict, org_funding_fallback: str = "") -> str:
    """يجمع التمويل من 3 حقول جديدة + fallback للحقل القديم + fallback للجمعية."""
    parts = []
    don = (init.get("funding_donations") or "").strip()
    par = (init.get("funding_partnerships") or "").strip()
    inst = (init.get("funding_institutional") or "").strip()
    if don and don != "غير مذكور":
        parts.append(f"💝 تبرعات: {don}")
    if par and par != "غير مذكور":
        parts.append(f"🤝 شراكات: {par}")
    if inst and inst != "غير مذكور":
        parts.append(f"🏛 مؤسسي: {inst}")
    if parts:
        return "\n".join(parts)
    # fallback للحقل القديم لو مازال موجود
    legacy = (init.get("funding_type") or "").strip()
    if legacy and legacy != "غير مذكور":
        return legacy
    # fallback لمصادر تمويل الجمعية ككل (من validation)
    if org_funding_fallback and org_funding_fallback.strip() and org_funding_fallback.strip() != "غير مذكور":
        return f"📍 على مستوى الجمعية: {org_funding_fallback.strip()}"
    return "—"


def _format_kpis(init: dict) -> str:
    """يحوّل قائمة الـ KPIs لنص مرتب — أو يستخدم الحقول القديمة كـ fallback."""
    kpis = init.get("kpis") or []
    if isinstance(kpis, list) and kpis:
        lines = []
        for k in kpis:
            if not isinstance(k, dict):
                continue
            ind = (k.get("indicator") or "").strip()
            val = (k.get("value") or "").strip()
            if ind and val:
                lines.append(f"• {ind}: {val}")
            elif ind:
                lines.append(f"• {ind}")
        if lines:
            return "\n".join(lines)
    # fallback للحقول القديمة
    metric = (init.get("metric") or "").strip()
    results = (init.get("results") or "").strip()
    pieces = [p for p in (metric, results) if p and p != "—"]
    return " | ".join(pieces) if pieces else "—"


def _build_matrix_sheet(ws, title, items, accepted_orgs=None):
    ws.merge_cells("A1:H1")
    ws["A1"] = title
    apply_style(ws["A1"], main_title_style())
    ws.row_dimensions[1].height = 40

    headers = [
        "اسم الجمعية", "اسم المبادرة", "وصف المبادرة",
        "الهدف الاستراتيجي",
        "الفئة المستهدفة (نظرياً)", "الفئة المستفيدة (فعلياً)",
        "مصادر التمويل", "مؤشرات الأداء",
    ]
    for col, h in enumerate(headers, 1):
        apply_style(ws.cell(row=3, column=col, value=h), header_style())
    ws.row_dimensions[3].height = 45

    # خريطة: اسم الجمعية → نص funding_sources من validation (للـ fallback)
    org_funding_map: dict[str, str] = {}
    for org in (accepted_orgs or []):
        oname = org.get("name", "")
        for f in org.get("fields", []):
            if f.get("name") == "funding_sources" and f.get("value"):
                org_funding_map[oname] = str(f.get("value"))
                break

    if not items:
        ws.merge_cells("A4:H4")
        ws["A4"] = "لا توجد مبادرات في هذا القسم"
        apply_style(ws["A4"], {
            "font": Font(name="IBM Plex Sans Arabic", size=11, italic=True, color="6B7280"),
            "fill": PatternFill("solid", fgColor=COLORS["cream"]),
            "alignment": Alignment(horizontal="center", vertical="center"),
        })
        ws.row_dimensions[4].height = 40

    for i, init in enumerate(items, 4):
        ws.cell(row=i, column=1, value=init.get("source_org", ""))
        ws.cell(row=i, column=2, value=init.get("name", ""))
        ws.cell(row=i, column=3, value=init.get("description") or "—")
        ws.cell(row=i, column=4, value=init.get("objective") or "—")
        # الفئة المستهدفة: target_group (جديد) → fallback لـ target_beneficiaries (قديم)
        target = (init.get("target_group") or init.get("target_beneficiaries") or "—").strip()
        ws.cell(row=i, column=5, value=target)
        # الفئة المستفيدة فعلياً: حقل جديد فقط
        ws.cell(row=i, column=6, value=(init.get("actual_beneficiaries") or "—").strip())
        org_fallback = org_funding_map.get(init.get("source_org", ""), "")
        ws.cell(row=i, column=7, value=_format_funding_summary(init, org_fallback))
        ws.cell(row=i, column=8, value=_format_kpis(init))

        for c in range(1, 9):
            apply_style(ws.cell(row=i, column=c), value_style())
        ws.row_dimensions[i].height = 90

    widths = [20, 22, 30, 22, 22, 22, 26, 26]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + col)].width = w


def _build_funding_sheet(ws, title, accepted, matrix):
    """جدول التمويل بالأنواع الـ 3 — مجمّع لكل جمعية من مبادراتها."""
    ws.merge_cells("A1:E1")
    ws["A1"] = title
    apply_style(ws["A1"], main_title_style())
    ws.row_dimensions[1].height = 40

    headers = ["الجمعية", "النوع", "💝 تبرعات", "🤝 شراكات", "🏛 تمويل مؤسسي"]
    for col, h in enumerate(headers, 1):
        apply_style(ws.cell(row=3, column=col, value=h), header_style())
    ws.row_dimensions[3].height = 32

    # نجمّع التمويل لكل جمعية من مبادراتها
    by_org: dict[str, dict[str, set]] = {}
    for org in accepted:
        oname = org.get("name", "")
        by_org[oname] = {"donations": set(), "partnerships": set(), "institutional": set(),
                         "type": "محلية" if org.get("type") == "local" else "دولية"}

    for init in matrix:
        oname = init.get("source_org", "")
        if oname not in by_org:
            continue
        for src_field, key in (
            ("funding_donations", "donations"),
            ("funding_partnerships", "partnerships"),
            ("funding_institutional", "institutional"),
        ):
            v = (init.get(src_field) or "").strip()
            if v and v != "غير مذكور":
                by_org[oname][key].add(v)
        # fallback لـ funding_type القديم → حطها في donations كـ best-effort
        legacy = (init.get("funding_type") or "").strip()
        if legacy and legacy != "غير مذكور" and not (
            init.get("funding_donations") or init.get("funding_partnerships") or init.get("funding_institutional")
        ):
            by_org[oname]["donations"].add(legacy)

    row = 4
    for oname, data in by_org.items():
        ws.cell(row=row, column=1, value=oname)
        ws.cell(row=row, column=2, value=data["type"])
        ws.cell(row=row, column=3, value="\n".join(f"• {x}" for x in sorted(data["donations"])) or "—")
        ws.cell(row=row, column=4, value="\n".join(f"• {x}" for x in sorted(data["partnerships"])) or "—")
        ws.cell(row=row, column=5, value="\n".join(f"• {x}" for x in sorted(data["institutional"])) or "—")
        for c in range(1, 6):
            apply_style(ws.cell(row=row, column=c), value_style())
        ws.row_dimensions[row].height = 95
        row += 1

    widths = [22, 12, 30, 30, 30]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + col)].width = w


def _build_kpis_sheet(ws, title, matrix):
    """جدول مؤشرات الأداء — صف لكل KPI."""
    ws.merge_cells("A1:D1")
    ws["A1"] = title
    apply_style(ws["A1"], main_title_style())
    ws.row_dimensions[1].height = 40

    headers = ["الجمعية", "المبادرة", "المؤشر", "القيمة"]
    for col, h in enumerate(headers, 1):
        apply_style(ws.cell(row=3, column=col, value=h), header_style())
    ws.row_dimensions[3].height = 32

    rows = []
    for init in matrix:
        oname = init.get("source_org", "")
        iname = init.get("name", "")
        kpis = init.get("kpis") or []
        if isinstance(kpis, list) and kpis:
            for k in kpis:
                if not isinstance(k, dict):
                    continue
                ind = (k.get("indicator") or "").strip()
                val = (k.get("value") or "").strip()
                if ind:
                    rows.append((oname, iname, ind, val or "—"))

    if not rows:
        ws.merge_cells("A4:D4")
        ws["A4"] = "لا توجد مؤشرات أداء مستخرجة"
        apply_style(ws["A4"], {
            "font": Font(name="IBM Plex Sans Arabic", size=11, italic=True, color="6B7280"),
            "fill": PatternFill("solid", fgColor=COLORS["cream"]),
            "alignment": Alignment(horizontal="center", vertical="center"),
        })
        ws.row_dimensions[4].height = 35
    else:
        for i, (oname, iname, ind, val) in enumerate(rows, 4):
            ws.cell(row=i, column=1, value=oname)
            ws.cell(row=i, column=2, value=iname)
            ws.cell(row=i, column=3, value=ind)
            ws.cell(row=i, column=4, value=val)
            for c in range(1, 5):
                apply_style(ws.cell(row=i, column=c), value_style())
            ws.row_dimensions[i].height = 30

    widths = [22, 25, 35, 22]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + col)].width = w


def generate_output_3(workflow_result, output_path):
    """مخرج 3: المصفوفة + قسم التمويل + قسم مؤشرات الأداء"""
    wb = Workbook()
    wb.remove(wb.active)

    matrix = workflow_result.get("initiative_matrix", [])
    local_matrix = [m for m in matrix if m.get("source_type") == "local"]
    intl_matrix = [m for m in matrix if m.get("source_type") == "international"]
    accepted = [o for o in workflow_result.get("local_organizations", []) +
                workflow_result.get("international_organizations", []) if o.get("is_accepted")]

    # ⭐ دايماً نعمل الـ 3 شيتس (حتى لو فاضي) عشان المستخدم يشوف التقسيم
    ws1 = wb.create_sheet("🇸🇦 المصفوفة المحلية")
    ws1.sheet_view.rightToLeft = True
    _build_matrix_sheet(ws1, f"المصفوفة المحلية ({len(local_matrix)} مبادرة)", local_matrix, accepted)

    ws2 = wb.create_sheet("🌍 المصفوفة الدولية")
    ws2.sheet_view.rightToLeft = True
    _build_matrix_sheet(ws2, f"المصفوفة الدولية ({len(intl_matrix)} مبادرة)", intl_matrix, accepted)

    ws3 = wb.create_sheet("⭐ المصفوفة الشاملة")
    ws3.sheet_view.rightToLeft = True
    _build_matrix_sheet(ws3, f"المصفوفة الشاملة — محلي + دولي ({len(matrix)} مبادرة)", matrix, accepted)

    # 💰 قسم التمويل + 📊 مؤشرات الأداء
    ws_fund = wb.create_sheet("💰 التمويل")
    ws_fund.sheet_view.rightToLeft = True
    _build_funding_sheet(ws_fund, "تصنيف التمويل: تبرعات / شراكات / تمويل مؤسسي", accepted, matrix)

    ws_kpi = wb.create_sheet("📊 مؤشرات الأداء")
    ws_kpi.sheet_view.rightToLeft = True
    _build_kpis_sheet(ws_kpi, "مؤشرات الأداء (KPIs)", matrix)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output))
    return str(output)


def generate_unified_report(workflow_result, output_path):
    """التقرير الموحد"""
    wb = Workbook()
    wb.remove(wb.active)
    
    client = workflow_result.get("client", {})
    accepted = [o for o in workflow_result.get("local_organizations", []) +
                workflow_result.get("international_organizations", []) if o.get("is_accepted")]
    matrix = workflow_result.get("initiative_matrix", [])
    
    # Summary
    ws = wb.create_sheet("📋 الملخص")
    ws.sheet_view.rightToLeft = True
    
    ws.merge_cells("A1:F1")
    ws["A1"] = "تقرير المقارنة المعيارية التنافسية — Unistep"
    apply_style(ws["A1"], main_title_style())
    ws.row_dimensions[1].height = 45
    
    ws.merge_cells("A2:F2")
    ws["A2"] = f"العميل: {client.get('name', '')}  |  {client.get('specialization', '')}  |  {datetime.now().strftime('%Y-%m-%d')}"
    apply_style(ws["A2"], {
        "font": Font(name="IBM Plex Sans Arabic", size=11, italic=True, color=COLORS["dark"]),
        "fill": PatternFill("solid", fgColor=COLORS["cream"]),
        "alignment": Alignment(horizontal="center", vertical="center"),
    })
    ws.row_dimensions[2].height = 28
    
    kpis = [
        ("جهات مقبولة", str(len(accepted))),
        ("مبادرات متطابقة", str(len(matrix))),
        ("جهات محلية", str(sum(1 for o in accepted if o.get("type") == "local"))),
        ("جهات دولية", str(sum(1 for o in accepted if o.get("type") == "international"))),
    ]
    row = 4
    for col, (label, val) in enumerate(kpis, 1):
        ws.cell(row=row, column=col*2-1, value=label)
        ws.cell(row=row+1, column=col*2-1, value=val)
        apply_style(ws.cell(row=row, column=col*2-1), {
            "font": Font(name="IBM Plex Sans Arabic", size=9),
            "fill": PatternFill("solid", fgColor=COLORS["amber_light"]),
            "alignment": Alignment(horizontal="center"),
        })
        apply_style(ws.cell(row=row+1, column=col*2-1), {
            "font": Font(name="IBM Plex Sans Arabic", size=20, bold=True, color=COLORS["dark"]),
            "fill": PatternFill("solid", fgColor=COLORS["amber_light"]),
            "alignment": Alignment(horizontal="center"),
        })
    ws.row_dimensions[row].height = 22
    ws.row_dimensions[row+1].height = 38
    for col in range(1, 9):
        ws.column_dimensions[chr(64 + col)].width = 18
    
    # المعايير
    ws_crit = wb.create_sheet("📋 معايير الاختيار")
    ws_crit.sheet_view.rightToLeft = True
    
    ws_crit.merge_cells("A1:C1")
    ws_crit["A1"] = "المخرج 1 — معايير اختيار المقارنات"
    apply_style(ws_crit["A1"], main_title_style())
    ws_crit.row_dimensions[1].height = 40
    
    for col, h in enumerate(["المستوى", "المعيار", "الوصف"], 1):
        apply_style(ws_crit.cell(row=3, column=col, value=h), header_style())
    ws_crit.row_dimensions[3].height = 30
    
    for i, (crit, desc, lvl, _) in enumerate(SELECTION_CRITERIA_7, 4):
        icon = "🔴" if lvl == "إلزامي" else "🟡"
        ws_crit.cell(row=i, column=1, value=f"{icon} {lvl}")
        ws_crit.cell(row=i, column=2, value=crit)
        ws_crit.cell(row=i, column=3, value=desc)
        for c in range(1, 4):
            apply_style(ws_crit.cell(row=i, column=c), value_style())
    
    ws_crit.column_dimensions["A"].width = 14
    ws_crit.column_dimensions["B"].width = 30
    ws_crit.column_dimensions["C"].width = 50
    
    # تقييم الجهات
    if accepted:
        ws_eval = wb.create_sheet("📊 تقييم الجهات")
        ws_eval.sheet_view.rightToLeft = True
        
        ws_eval.merge_cells("A1:H1")
        ws_eval["A1"] = "تقييم الجهات على المعايير"
        apply_style(ws_eval["A1"], main_title_style())
        ws_eval.row_dimensions[1].height = 40
        
        headers = ["الجهة"] + [c[0][:14] for c in SELECTION_CRITERIA_7]
        for col, h in enumerate(headers, 1):
            apply_style(ws_eval.cell(row=3, column=col, value=h), header_style())
        ws_eval.row_dimensions[3].height = 50
        
        for i, org in enumerate(accepted, 4):
            ws_eval.cell(row=i, column=1, value=org.get("name", ""))
            for j, (_, _, _, cid) in enumerate(SELECTION_CRITERIA_7, 2):
                val = org.get(cid, False)
                ws_eval.cell(row=i, column=j, value="✅" if val else "—")
            for c in range(1, len(headers) + 1):
                apply_style(ws_eval.cell(row=i, column=c), value_style())
        
        ws_eval.column_dimensions["A"].width = 25
        for col in range(2, 9):
            ws_eval.column_dimensions[chr(64 + col)].width = 14
    
    # Sheets لكل جهة
    for i, org in enumerate(accepted, 1):
        sheet_name = f"{i}. {org.get('name', '')[:25]}"
        ws = wb.create_sheet(sheet_name[:31])
        ws.sheet_view.rightToLeft = True
        
        ws.merge_cells("A1:B1")
        ws["A1"] = org.get("name", "")
        apply_style(ws["A1"], main_title_style())
        ws.row_dimensions[1].height = 40
        
        ws.merge_cells("A2:B2")
        ws["A2"] = f"{org.get('url', '')}  |  {org.get('country', '')}"
        apply_style(ws["A2"], {
            "font": Font(name="IBM Plex Sans Arabic", size=9, italic=True),
            "fill": PatternFill("solid", fgColor=COLORS["cream"]),
            "alignment": Alignment(horizontal="center", vertical="center"),
        })
        
        apply_style(ws.cell(row=4, column=1, value="اسم البيان"), header_style())
        apply_style(ws.cell(row=4, column=2, value="التفاصيل"), header_style())
        ws.row_dimensions[4].height = 30
        
        fields_dict = {f.get("name"): f for f in org.get("fields", [])}
        row = 5
        for fname in ROW_ORDER:
            label = FIELD_LABELS.get(fname, fname)
            f = fields_dict.get(fname)
            value = f.get("value") if f and f.get("value") else "غير متاح"
            
            apply_style(ws.cell(row=row, column=1, value=label), label_style())
            apply_style(ws.cell(row=row, column=2, value=str(value)), value_style())
            ws.row_dimensions[row].height = min(180, max(35, len(str(value)) // 5))
            row += 1
        
        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 75
    
    # المصفوفات (3 شيتس: محلي / دولي / شامل)
    local_matrix = [m for m in matrix if m.get("source_type") == "local"]
    intl_matrix = [m for m in matrix if m.get("source_type") == "international"]

    ws_local = wb.create_sheet("🇸🇦 المصفوفة المحلية")
    ws_local.sheet_view.rightToLeft = True
    _build_matrix_sheet(ws_local, f"المصفوفة المحلية ({len(local_matrix)} مبادرة)", local_matrix, accepted)

    ws_intl = wb.create_sheet("🌍 المصفوفة الدولية")
    ws_intl.sheet_view.rightToLeft = True
    _build_matrix_sheet(ws_intl, f"المصفوفة الدولية ({len(intl_matrix)} مبادرة)", intl_matrix, accepted)

    ws_m = wb.create_sheet("⭐ المصفوفة الشاملة")
    ws_m.sheet_view.rightToLeft = True
    _build_matrix_sheet(ws_m, f"المصفوفة الشاملة — محلي + دولي ({len(matrix)} مبادرة)", matrix, accepted)

    # 💰 قسم التمويل + 📊 مؤشرات الأداء
    ws_fund = wb.create_sheet("💰 التمويل")
    ws_fund.sheet_view.rightToLeft = True
    _build_funding_sheet(ws_fund, "تصنيف التمويل: تبرعات / شراكات / تمويل مؤسسي", accepted, matrix)

    ws_kpi = wb.create_sheet("📊 مؤشرات الأداء")
    ws_kpi.sheet_view.rightToLeft = True
    _build_kpis_sheet(ws_kpi, "مؤشرات الأداء (KPIs)", matrix)

    # Audit
    ws_audit = wb.create_sheet("🔒 سجل التحقق")
    ws_audit.sheet_view.rightToLeft = True
    
    ws_audit.merge_cells("A1:E1")
    ws_audit["A1"] = "سجل التحقق — كل بيان مع مصدره الحرفي"
    apply_style(ws_audit["A1"], main_title_style())
    ws_audit.row_dimensions[1].height = 40
    
    for col, h in enumerate(["الجمعية", "البيان", "القيمة", "الاقتباس الحرفي", "المصدر"], 1):
        apply_style(ws_audit.cell(row=3, column=col, value=h), header_style())
    ws_audit.row_dimensions[3].height = 30
    
    row = 4
    for org in accepted:
        for f in org.get("fields", []):
            if not f.get("passed"):
                continue
            label = FIELD_LABELS.get(f.get("name", ""), f.get("name", ""))
            ws_audit.cell(row=row, column=1, value=org.get("name", ""))
            ws_audit.cell(row=row, column=2, value=label)
            ws_audit.cell(row=row, column=3, value=str(f.get("value", ""))[:300])
            ws_audit.cell(row=row, column=4, value=str(f.get("exact_quote", ""))[:500])
            ws_audit.cell(row=row, column=5, value=org.get("url", ""))
            for c in range(1, 6):
                apply_style(ws_audit.cell(row=row, column=c), value_style())
            ws_audit.row_dimensions[row].height = 80
            row += 1
    
    widths = [22, 18, 35, 50, 25]
    for col, w in enumerate(widths, 1):
        ws_audit.column_dimensions[chr(64 + col)].width = w
    
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output))
    return str(output)
