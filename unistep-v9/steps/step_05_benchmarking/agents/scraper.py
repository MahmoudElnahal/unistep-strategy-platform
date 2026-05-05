"""
═══════════════════════════════════════════════════════════
Enhanced Web Scraper v7
═══════════════════════════════════════════════════════════

- requests أولاً (سريع)
- Selenium fallback (للمواقع الديناميكية)
- يكتشف PDFs ويحملها
- يحفظ raw text + raw HTML
"""
import io
import logging
import re
import time
from typing import Optional
from dataclasses import dataclass, field
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup

try:
    from pypdf import PdfReader
    HAS_PYPDF = True
except ImportError:
    HAS_PYPDF = False

try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

logger = logging.getLogger(__name__)


ANNUAL_REPORT_KEYWORDS = [
    "تقرير سنوي", "التقرير السنوي", "تقارير سنوية", "تقريرنا السنوي",
    "تقرير الإنجازات", "تقرير سنويه",
    "annual report", "annual-report", "annualreport",
    "yearly report", "year in review",
    "2020", "2021", "2022", "2023", "2024", "2025",
]


@dataclass
class DocumentLink:
    url: str
    doc_type: str  # 'pdf' | 'docx' | 'doc'
    anchor_text: str = ""
    is_annual_report: bool = False


@dataclass
class ScrapedPage:
    url: str
    title: str
    text: str
    html: str
    status_code: int
    discovered_links: list[str] = field(default_factory=list)
    link_anchors: dict[str, str] = field(default_factory=dict)
    pdf_links: list[str] = field(default_factory=list)
    document_links: list[DocumentLink] = field(default_factory=list)
    error: Optional[str] = None
    method_used: str = "requests"

    @property
    def is_success(self) -> bool:
        return self.status_code == 200 and self.error is None and len(self.text) > 100

    @property
    def text_length(self) -> int:
        return len(self.text)


class EnhancedScraper:
    USER_AGENT = (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/131.0.0.0 Safari/537.36"
    )

    HEADERS = {
        "User-Agent": USER_AGENT,
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "ar,en;q=0.9",
        "Accept-Encoding": "gzip, deflate, br",
        "Connection": "keep-alive",
    }

    LINK_KEYWORDS = {
        "about": ["من-نحن", "من نحن", "عن", "نبذة", "about", "who-we-are"],
        "programs": ["برامج", "مبادرات", "خدمات", "أنشطة", "programs", "initiatives", "services"],
        "reports": ["تقارير", "إنجازات", "إفصاح", "reports", "publications", "achievements", "annual"],
        "vision": ["رؤية", "رسالة", "أهداف", "vision", "mission", "goals", "strategy"],
        "donate": ["تبرع", "تبرعات", "ادعم", "donate", "support", "resources", "موارد"],
        "media": ["إعلام", "أخبار", "media", "news", "press"],
    }

    def __init__(self, use_selenium: bool = True, timeout: int = 25):
        self.use_selenium = use_selenium
        self.timeout = timeout
        self._driver = None

    def _get_selenium_driver(self):
        if self._driver is not None:
            return self._driver

        try:
            from selenium import webdriver
            from selenium.webdriver.chrome.options import Options

            options = Options()
            options.add_argument("--headless=new")
            options.add_argument("--no-sandbox")
            options.add_argument("--disable-dev-shm-usage")
            options.add_argument("--disable-blink-features=AutomationControlled")
            options.add_argument(f"--user-agent={self.USER_AGENT}")
            options.add_argument("--lang=ar")
            options.add_argument("--window-size=1920,1080")
            options.add_experimental_option("excludeSwitches", ["enable-automation"])
            options.add_experimental_option("useAutomationExtension", False)

            self._driver = webdriver.Chrome(options=options)
            self._driver.set_page_load_timeout(self.timeout)
            self._driver.execute_script(
                "Object.defineProperty(navigator, 'webdriver', {get: () => undefined})"
            )
            logger.info("✅ Selenium driver جاهز")
            return self._driver

        except Exception as e:
            logger.warning(f"⚠️ Selenium failed: {e}")
            self.use_selenium = False
            return None

    def close(self):
        if self._driver is not None:
            try:
                self._driver.quit()
            except Exception:
                pass
            self._driver = None

    def fetch(self, url: str) -> ScrapedPage:
        logger.info(f"🌐 جلب: {url}")

        # 1. requests (verify=True)
        page = self._fetch_with_requests(url, verify_ssl=True)
        if page.is_success and page.text_length >= 500:
            logger.info(f"  ✅ requests نجح ({page.text_length} حرف)")
            return page

        # 2. requests مع SSL متساهل (لمواقع .sa اللي شهاداتها بايظة)
        if not page.is_success or page.text_length < 500:
            logger.info(f"  🔄 retry بـ verify=False...")
            page2 = self._fetch_with_requests(url, verify_ssl=False)
            if page2.is_success and page2.text_length >= 500:
                logger.info(f"  ✅ requests (no-ssl) نجح ({page2.text_length} حرف)")
                return page2
            if page2.text_length > page.text_length:
                page = page2

        # 3. Selenium fallback
        if self.use_selenium:
            logger.info(f"  🔄 ينتقل لـ Selenium...")
            sp = self._fetch_with_selenium(url)
            if sp.is_success:
                return sp
            if sp.text_length > page.text_length:
                page = sp

        if not page.is_success:
            logger.warning(f"  ❌ كل المحاولات فشلت لـ {url} — السبب: {page.error or 'نص قصير'}")
        return page

    def _fetch_with_requests(self, url: str, verify_ssl: bool = True) -> ScrapedPage:
        try:
            if not verify_ssl:
                import urllib3
                urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
            r = requests.get(url, headers=self.HEADERS, timeout=self.timeout,
                            allow_redirects=True, verify=verify_ssl)
            soup = BeautifulSoup(r.content, "lxml")

            for tag in soup(["script", "style", "noscript", "iframe"]):
                tag.decompose()

            title = (soup.title.string if soup.title else "") or ""
            text = soup.get_text(separator="\n", strip=True)
            text = "\n".join(line for line in text.split("\n") if line.strip())

            important, anchors, _ = self._extract_links(soup, url)
            documents = self._find_documents(soup, url)

            return ScrapedPage(
                url=url, title=title.strip(), text=text,
                html=str(soup)[:80000], status_code=r.status_code,
                discovered_links=important,
                link_anchors=anchors,
                pdf_links=[d.url for d in documents if d.doc_type == "pdf"],
                document_links=documents,
                method_used="requests" if verify_ssl else "requests-nossl",
            )

        except Exception as e:
            logger.warning(f"  ⚠️ requests({verify_ssl=}) فشل: {type(e).__name__}: {e}")
            return ScrapedPage(url=url, title="", text="", html="",
                              status_code=0, error=str(e),
                              method_used="requests" if verify_ssl else "requests-nossl")

    def _fetch_with_selenium(self, url: str) -> ScrapedPage:
        driver = self._get_selenium_driver()
        if driver is None:
            return ScrapedPage(url=url, title="", text="", html="",
                              status_code=0, error="Selenium unavailable",
                              method_used="selenium")

        try:
            driver.get(url)
            time.sleep(5)  # وقت أطول للـ SPA

            # Scroll لتحريك lazy loading
            try:
                driver.execute_script("window.scrollTo(0, document.body.scrollHeight/2);")
                time.sleep(1)
                driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
                time.sleep(2)
                driver.execute_script("window.scrollTo(0, 0);")
                time.sleep(1)
            except Exception:
                pass

            html = driver.page_source
            title = driver.title or ""
            soup = BeautifulSoup(html, "lxml")

            for tag in soup(["script", "style", "noscript", "iframe"]):
                tag.decompose()

            text = soup.get_text(separator="\n", strip=True)
            text = "\n".join(line for line in text.split("\n") if line.strip())

            important, anchors, _ = self._extract_links(soup, url)
            documents = self._find_documents(soup, url)

            return ScrapedPage(
                url=url, title=title.strip(), text=text,
                html=html[:80000], status_code=200,
                discovered_links=important,
                link_anchors=anchors,
                pdf_links=[d.url for d in documents if d.doc_type == "pdf"],
                document_links=documents,
                method_used="selenium",
            )

        except Exception as e:
            logger.error(f"❌ Selenium فشل لـ {url}: {type(e).__name__}: {e}")
            return ScrapedPage(url=url, title="", text="", html="",
                              status_code=0, error=str(e), method_used="selenium")

    def _extract_links(self, soup: BeautifulSoup, base_url: str) -> tuple[list[str], dict[str, str], list[tuple[str, str]]]:
        """
        Returns (important_links, link_anchor_map, all_internal_links_with_text).
        - important_links: links matching about/programs/reports keywords (legacy behavior)
        - link_anchor_map: {url: anchor_text} for ALL internal links (used for initiative name matching)
        - all_internal_links_with_text: every internal anchor with its text — used by deep-crawl
        """
        important = set()
        anchor_map: dict[str, str] = {}
        all_internal: list[tuple[str, str]] = []
        base_domain = urlparse(base_url).netloc

        for link in soup.find_all("a", href=True):
            href = link.get("href", "")
            text = link.get_text(strip=True)

            if not href or href.startswith("#") or href.startswith("javascript:"):
                continue
            if href.lower().endswith((".pdf", ".docx", ".doc", ".xlsx", ".zip", ".jpg", ".png")):
                continue

            full_url = urljoin(base_url, href)
            if urlparse(full_url).netloc != base_domain:
                continue

            text_lower = text.lower()
            href_lower = href.lower()

            if text:
                anchor_map[full_url] = text
                all_internal.append((full_url, text))

            for category, keywords in self.LINK_KEYWORDS.items():
                if any(kw in text_lower or kw in href_lower for kw in keywords):
                    important.add(full_url)
                    break

        return list(important)[:25], anchor_map, all_internal

    def _find_documents(self, soup: BeautifulSoup, base_url: str) -> list[DocumentLink]:
        docs: dict[str, DocumentLink] = {}

        for link in soup.find_all("a", href=True):
            href = link.get("href", "")
            text = link.get_text(strip=True)
            text_lower = text.lower()
            href_lower = href.lower()

            doc_type = None
            if href_lower.endswith(".pdf") or ".pdf?" in href_lower or ".pdf#" in href_lower:
                doc_type = "pdf"
            elif href_lower.endswith(".docx") or ".docx?" in href_lower:
                doc_type = "docx"
            elif href_lower.endswith(".doc") or ".doc?" in href_lower:
                doc_type = "doc"

            if not doc_type:
                continue

            full = urljoin(base_url, href)
            if full in docs:
                continue

            haystack = f"{text_lower} {href_lower}"
            is_annual = any(kw.lower() in haystack for kw in ANNUAL_REPORT_KEYWORDS)

            docs[full] = DocumentLink(
                url=full,
                doc_type=doc_type,
                anchor_text=text,
                is_annual_report=is_annual,
            )

        return list(docs.values())

    def fetch_docx(self, url: str, max_chars: int = 200000) -> ScrapedPage:
        """تنزيل DOCX واستخراج النص — للتقارير السنوية بصيغة Word."""
        logger.info(f"📝 DOCX: {url}")
        if not HAS_DOCX:
            return ScrapedPage(url=url, title="", text="", html="",
                              status_code=0, error="python-docx not installed",
                              method_used="docx")
        try:
            r = requests.get(url, headers=self.HEADERS, timeout=self.timeout,
                            allow_redirects=True)
            if r.status_code != 200:
                return ScrapedPage(url=url, title="", text="", html="",
                                  status_code=r.status_code, error=f"http {r.status_code}",
                                  method_used="docx")
            content = r.content
            if len(content) > 25 * 1024 * 1024:
                logger.warning(f"  ⚠️ DOCX كبير جداً ({len(content)//1024//1024}MB) — تخطي")
                return ScrapedPage(url=url, title="", text="", html="",
                                  status_code=200, error="docx too large",
                                  method_used="docx")

            doc = DocxDocument(io.BytesIO(content))
            parts: list[str] = []
            for p in doc.paragraphs:
                if p.text.strip():
                    parts.append(p.text)
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            text = "\n".join(parts)[:max_chars]
            logger.info(f"  ✅ DOCX text: {len(text)} حرف")
            return ScrapedPage(
                url=url, title=url.rsplit("/", 1)[-1], text=text,
                html="", status_code=200, method_used="docx",
            )
        except Exception as e:
            logger.warning(f"  ⚠️ DOCX فشل: {e}")
            return ScrapedPage(url=url, title="", text="", html="",
                              status_code=0, error=str(e), method_used="docx")

    def fetch_document(self, doc_link: DocumentLink) -> ScrapedPage:
        """يوزع على fetch_pdf أو fetch_docx حسب نوع الملف."""
        if doc_link.doc_type == "pdf":
            return self.fetch_pdf(doc_link.url)
        if doc_link.doc_type == "docx":
            return self.fetch_docx(doc_link.url)
        return ScrapedPage(url=doc_link.url, title="", text="", html="",
                          status_code=0, error=f"unsupported doc type: {doc_link.doc_type}",
                          method_used="document")

    def fetch_pdf(self, url: str, max_pages: int = 30) -> ScrapedPage:
        """تنزيل PDF واستخراج النص (للتقارير السنوية)"""
        logger.info(f"📄 PDF: {url}")
        if not HAS_PYPDF:
            return ScrapedPage(url=url, title="", text="", html="",
                              status_code=0, error="pypdf not installed",
                              method_used="pdf")
        try:
            r = requests.get(url, headers=self.HEADERS, timeout=self.timeout,
                            allow_redirects=True, stream=True)
            if r.status_code != 200:
                return ScrapedPage(url=url, title="", text="", html="",
                                  status_code=r.status_code, error=f"http {r.status_code}",
                                  method_used="pdf")
            content = r.content
            if len(content) > 25 * 1024 * 1024:  # 25 MB cap
                logger.warning(f"  ⚠️ PDF كبير جداً ({len(content)//1024//1024}MB) — تخطي")
                return ScrapedPage(url=url, title="", text="", html="",
                                  status_code=200, error="pdf too large",
                                  method_used="pdf")

            reader = PdfReader(io.BytesIO(content))
            parts = []
            for i, page in enumerate(reader.pages[:max_pages]):
                try:
                    txt = page.extract_text() or ""
                    if txt.strip():
                        parts.append(txt)
                except Exception:
                    continue
            text = "\n".join(parts)
            text = "\n".join(line for line in text.split("\n") if line.strip())
            logger.info(f"  ✅ PDF text: {len(text)} حرف من {len(reader.pages)} صفحة")
            return ScrapedPage(
                url=url, title=url.rsplit("/", 1)[-1], text=text,
                html="", status_code=200, method_used="pdf",
            )
        except Exception as e:
            logger.warning(f"  ⚠️ PDF فشل: {e}")
            return ScrapedPage(url=url, title="", text="", html="",
                              status_code=0, error=str(e), method_used="pdf")

    def fetch_organization_pages(self, base_url: str) -> dict[str, ScrapedPage]:
        """جلب أهم 5-7 صفحات للجهة + المستندات (تقارير سنوية PDF/DOCX)."""
        results = {}
        home = self.fetch(base_url)
        results["home"] = home

        if not home.is_success:
            return results

        seen_types = set()
        all_documents: dict[str, DocumentLink] = {d.url: d for d in home.document_links}

        for link in home.discovered_links[:18]:
            text_lower = link.lower()

            page_type = None
            if any(kw in text_lower for kw in ["about", "من-نحن", "نبذة", "عن"]):
                page_type = "about"
            elif any(kw in text_lower for kw in ["program", "برامج", "مبادر", "خدمات"]):
                page_type = "programs"
            elif any(kw in text_lower for kw in ["report", "تقرير", "إنجاز", "annual"]):
                page_type = "reports"
            elif any(kw in text_lower for kw in ["vision", "رؤية", "رسالة", "أهداف"]):
                page_type = "vision"
            elif any(kw in text_lower for kw in ["donate", "تبرع", "ادعم", "support"]):
                page_type = "donate"
            elif any(kw in text_lower for kw in ["news", "media", "أخبار", "إعلام"]):
                page_type = "media"

            if page_type and page_type not in seen_types:
                page = self.fetch(link)
                results[page_type] = page
                seen_types.add(page_type)
                for d in page.document_links:
                    all_documents.setdefault(d.url, d)
                if len(seen_types) >= 6:
                    break

        # ⭐ تنزيل المستندات: نقدّم التقارير السنوية أولاً (DOCX قبل PDF لو نفس التقرير)
        docs_sorted = sorted(
            all_documents.values(),
            key=lambda d: (
                0 if d.is_annual_report else 1,
                0 if d.doc_type == "docx" else 1,
            ),
        )

        annual_count = 0
        other_count = 0
        for doc in docs_sorted:
            if annual_count >= 3 and other_count >= 2:
                break
            if doc.is_annual_report and annual_count >= 3:
                continue
            if not doc.is_annual_report and other_count >= 2:
                continue

            page = self.fetch_document(doc)
            if not page.is_success:
                continue

            if doc.is_annual_report:
                results[f"annual_report_{annual_count}"] = page
                annual_count += 1
            else:
                results[f"document_{other_count}"] = page
                other_count += 1

        if annual_count or other_count:
            logger.info(
                f"📄 تم تحميل {annual_count} تقرير سنوي + {other_count} مستند آخر"
            )

        return results

    @staticmethod
    def _normalize_arabic(text: str) -> str:
        """يبسّط نص عربي للمقارنة: يشيل التشكيل، توحيد الألف والياء، lowercase."""
        if not text:
            return ""
        # شيل التشكيل
        text = re.sub(r"[ً-ْٰـ]", "", text)
        # توحيد الألف
        text = re.sub(r"[إأآا]", "ا", text)
        # توحيد الياء
        text = text.replace("ى", "ي").replace("ئ", "ي")
        # توحيد التاء المربوطة
        text = text.replace("ة", "ه")
        # شيل الترقيم
        text = re.sub(r"[^\w\s]", " ", text)
        text = re.sub(r"\s+", " ", text).strip().lower()
        return text



_scraper_instance: Optional[EnhancedScraper] = None


def get_scraper() -> EnhancedScraper:
    global _scraper_instance
    if _scraper_instance is None:
        _scraper_instance = EnhancedScraper(use_selenium=True)
    return _scraper_instance


def cleanup_scraper():
    global _scraper_instance
    if _scraper_instance is not None:
        _scraper_instance.close()
        _scraper_instance = None
